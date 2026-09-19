"""公文格式规范 / 批量套打等「文档生成」类能力。

与其它模块的区别：本模块的产出是 **Word 文档或批量 PDF**，而不是对单个 PDF 原地处理。

排版依据：《党政机关公文格式》GB/T 9704-2012
  - 页面：A4，页边距 上 37mm / 下 35mm / 左 28mm / 右 26mm
  - 标题：方正小标宋简体 二号（22pt）居中
  - 正文：仿宋_GB2312 三号（16pt），固定行距 28 磅，首行缩进 2 字符
  - 层级：一、→ 黑体三号 ； （一）→ 楷体三号 ； 1. → 仿宋三号加粗

⚠️ 字体说明：上述为公文字体，属**商业字体**，不能随项目分发。
   若使用者机器未安装，Word 会回退到相近字体（排版结构仍然正确）。
"""
import re
from pathlib import Path

from .core import register, save_uploads, new_tmp, send_file

# ---------- 字号与行距（磅） ----------
SZ_ERHAO = 22      # 二号：标题
SZ_SANHAO = 16     # 三号：正文 / 层级标题
LINE_EXACT = 28    # 正文固定行距 28 磅

FONT_TITLE = "方正小标宋简体"
FONT_BODY = "仿宋_GB2312"
FONT_H1 = "黑体"
FONT_H2 = "楷体_GB2312"

# 正文层级识别
RE_H1 = re.compile(r"^[一二三四五六七八九十百零]+、")                 # 一、
RE_H2 = re.compile(r"^[（(][一二三四五六七八九十百零]+[）)]")          # （一）
RE_H3 = re.compile(r"^\d+\s*[．.、]")                                # 1. / 1、


def _set_font(run, cn_font, size_pt, bold=False):
    """设置字体：必须同时写 w:eastAsia，否则 Word 打开后会回退成默认字体。"""
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = cn_font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), cn_font)
    # 西文与数字用 Times New Roman（公文惯例）
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")


def _setup_page(doc):
    """A4 + 党政机关公文页边距。"""
    from docx.shared import Cm

    s = doc.sections[0]
    s.page_width = Cm(21)
    s.page_height = Cm(29.7)
    s.top_margin = Cm(3.7)
    s.bottom_margin = Cm(3.5)
    s.left_margin = Cm(2.8)
    s.right_margin = Cm(2.6)


def _add_para(doc, text, font, size, *, center=False, right=False,
              indent_chars=0, bold=False, space_after=0):
    """追加一个规范段落（固定行距 + 首行缩进 + 指定中文字体）。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.shared import Pt

    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(LINE_EXACT)
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif right:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if indent_chars:
        pf.first_line_indent = Pt(size * indent_chars)
    _set_font(p.add_run(text), font, size, bold)
    return p


def _lines_from_upload(files, params):
    """取正文行：优先 Word 文件，其次 txt/md 文件，最后取表单里的 text。"""
    text = (params.get("text") or "").strip()
    if files:
        saved = save_uploads(files)
        docx_saved = next((s for s in saved if s[2].lower() == ".docx"), None)
        if docx_saved:
            from docx import Document
            return [p.text for p in Document(str(docx_saved[0])).paragraphs]
        txt_saved = next((s for s in saved if s[2].lower() in (".txt", ".md")), None)
        if txt_saved:
            return Path(txt_saved[0]).read_text(encoding="utf-8", errors="ignore").splitlines()
    return text.splitlines() if text else []


@register("gongwen-format", needs_files=False, desc="公文格式规范")
def gongwen_format(files, params):
    """把上传的 Word / 粘贴的正文，按 GB/T 9704-2012 重新排版。

    返回 **JSON** 而不是直接给文件：包含「处理前原文 + 处理后版式结构 + 文档 base64」，
    让前端可以「先预览、确认无误再下载」，而不是生成即下载。
    """
    import base64

    from docx import Document

    raw_lines = [ln.rstrip() for ln in _lines_from_upload(files, params)]
    while raw_lines and not raw_lines[0].strip():
        raw_lines.pop(0)
    while raw_lines and not raw_lines[-1].strip():
        raw_lines.pop()
    if not raw_lines:
        return {"ok": False, "level": "warn", "message": "请粘贴公文正文，或上传 .docx / .txt 文件。"}

    title = (params.get("title") or "").strip()
    body = raw_lines
    if not title:
        title = body[0].strip()          # 约定：首行作为标题
        body = body[1:]
    if not title:
        return {"ok": False, "level": "warn", "message": "未识别到标题，请填写标题或把标题放在首行。"}

    doc = Document()
    _setup_page(doc)
    blocks = []

    def emit(text, font, size, *, kind="body", center=False, right=False,
             indent=0, bold=False, extra_after=0):
        """写入 Word 的同时记录一条预览块，保证「所见 = 下载所得」。"""
        _add_para(doc, text, font, size, center=center, right=right,
                  indent_chars=indent, bold=bold, space_after=extra_after)
        blocks.append({
            "kind": kind,
            "text": text,
            "font": font,
            "sizePt": size,
            "bold": bold,
            "align": "center" if center else ("right" if right else "left"),
            "indentChars": indent,
        })

    # 标题：小标宋 二号 居中
    emit(title, FONT_TITLE, SZ_ERHAO, kind="title", center=True, extra_after=10)

    # 主送机关（顶格）
    receiver = (params.get("receiver") or "").strip()
    if receiver:
        emit(receiver, FONT_BODY, SZ_SANHAO, kind="receiver")

    # 正文：按层级套字体
    for line in body:
        t = line.strip()
        if not t:
            continue
        if RE_H1.match(t):
            emit(t, FONT_H1, SZ_SANHAO, kind="h1")
        elif RE_H2.match(t):
            emit(t, FONT_H2, SZ_SANHAO, kind="h2")
        elif RE_H3.match(t):
            emit(t, FONT_BODY, SZ_SANHAO, kind="h3", bold=True)
        else:
            emit(t, FONT_BODY, SZ_SANHAO, kind="body", indent=2)

    # 落款（单位 / 日期，右对齐）
    for key in ("signer", "date"):
        val = (params.get(key) or "").strip()
        if val:
            emit(val, FONT_BODY, SZ_SANHAO, kind=key, right=True)

    out = new_tmp() / "gongwen.docx"
    doc.save(str(out))
    data = out.read_bytes()

    return {
        "ok": True,
        "preview": {
            "page": {
                "widthCm": 21, "heightCm": 29.7,
                "marginTopCm": 3.7, "marginBottomCm": 3.5,
                "marginLeftCm": 2.8, "marginRightCm": 2.6,
                "lineSpacingPt": LINE_EXACT,
            },
            # 处理前：原始行（未套用任何格式）
            "before": [ln.strip() for ln in raw_lines if ln.strip()],
            # 处理后：逐段版式（前端据此渲染 A4 预览）
            "blocks": blocks,
        },
        "file": {
            "name": f"{title[:30] or '公文'}.docx",
            "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "size": len(data),
            "base64": base64.b64encode(data).decode("ascii"),
        },
    }
